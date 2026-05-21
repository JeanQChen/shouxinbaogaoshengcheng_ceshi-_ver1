"""Markdown → docx export.

Line-by-line state-machine parser. Supports:
  - H1-H4 headings
  - Paragraphs with inline **bold** and *italic*
  - Pipe tables
  - Unordered lists (- item)
  - Ordered lists (1. item)
  - Colored <span> tags (from verifier annotations)
"""

import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Styling constants ──

FONT_BODY = "SimSun"
FONT_HEADING = "SimHei"
FONT_SIZE_BODY = Pt(12)
FONT_SIZE_H1 = Pt(18)
FONT_SIZE_H2 = Pt(16)
FONT_SIZE_H3 = Pt(14)
FONT_SIZE_H4 = Pt(12)
LINE_SPACING = 1.5
MARGIN = Cm(2.5)

_COLOR_SPAN_RE = re.compile(
    r'<span\s+style="([^"]*)"[^>]*>(.*?)</span>',
    re.DOTALL,
)
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_ITALIC_RE = re.compile(r"\*(.+?)\*")

# ── Main entry ──


def export(markdown: str, output_path: str) -> None:
    """Export Markdown report to Word document and write to output_path."""
    doc = Document()

    # Page setup
    for section in doc.sections:
        section.top_margin = MARGIN
        section.bottom_margin = MARGIN
        section.left_margin = MARGIN
        section.right_margin = MARGIN

    style = doc.styles["Normal"]
    font = style.font
    font.name = FONT_BODY
    font.size = FONT_SIZE_BODY
    pf = style.paragraph_format
    pf.line_spacing = LINE_SPACING

    lines = markdown.split("\n")
    table_buffer: list[str] = []
    para_buffer: list[str] = []
    in_table = False
    pending_header: str | None = None  # potential table header line

    def _flush_para():
        nonlocal para_buffer, pending_header
        if para_buffer:
            text = " ".join(para_buffer).strip()
            para_buffer.clear()
            if text:
                _add_paragraph(doc, text)
        pending_header = None

    def _flush_table():
        nonlocal table_buffer, in_table, pending_header
        if table_buffer:
            _add_table(doc, table_buffer)
            table_buffer.clear()
        in_table = False
        pending_header = None

    for line in lines:
        stripped = line.strip()

        # Table: separator line — previous line might be header
        if re.match(r"^\|[-:\s|]+\|\s*$", stripped):
            if pending_header is not None:
                table_buffer.append(pending_header)
                pending_header = None
            _flush_para()
            in_table = True
            continue

        # Potential table header/data line (starts and ends with |)
        if stripped.startswith("|") and stripped.endswith("|") and "---" not in stripped:
            if in_table:
                table_buffer.append(stripped)
                continue
            elif pending_header is None:
                # Might be a header — delay and check next line
                pending_header = stripped
                continue

        # Not a table line — if we had a pending header it was just a paragraph
        if pending_header is not None:
            para_buffer.append(pending_header)
            pending_header = None

        if in_table and not stripped.startswith("|"):
            _flush_table()

        # Empty line
        if not stripped:
            _flush_table()
            _flush_para()
            continue

        # Heading
        heading_match = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading_match:
            _flush_table()
            _flush_para()
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()
            _add_heading(doc, text, level)
            continue

        # Horizontal rule
        if stripped in ("---", "***", "___", "* * *"):
            _flush_table()
            _flush_para()
            doc.add_paragraph("_" * 60)
            continue

        # Unordered list
        list_match = re.match(r"^-\s+(.+)$", stripped)
        if list_match:
            _flush_table()
            _flush_para()
            text = list_match.group(1)
            p = doc.add_paragraph(style="List Bullet")
            _add_inline_runs(p, text)
            continue

        # Ordered list
        ol_match = re.match(r"^\d+\.\s+(.+)$", stripped)
        if ol_match:
            _flush_table()
            _flush_para()
            text = ol_match.group(1)
            p = doc.add_paragraph(style="List Number")
            _add_inline_runs(p, text)
            continue

        # Normal paragraph
        para_buffer.append(stripped)

    # End-of-file flush
    if pending_header is not None:
        para_buffer.append(pending_header)
        pending_header = None
    _flush_table()
    _flush_para()

    doc.save(output_path)


# ── Element builders ──


def _add_heading(doc: Document, text: str, level: int) -> None:
    sizes = {1: FONT_SIZE_H1, 2: FONT_SIZE_H2, 3: FONT_SIZE_H3, 4: FONT_SIZE_H4}
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    run.font.name = FONT_HEADING
    run.font.size = sizes.get(level, FONT_SIZE_H4)
    run.bold = True
    if level == 1:
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _add_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    _add_inline_runs(p, text)


def _add_inline_runs(p, text: str) -> None:
    """Process inline formatting: colored spans, bold, italic."""
    parts = _tokenize_spans(text)
    for part_text, color, bg_color in parts:
        _add_text_with_formatting(p, part_text, color, bg_color)


def _tokenize_spans(text: str) -> list[tuple[str, str | None, str | None]]:
    """Split text into (text, text_color, bg_color) segments from <span> tags."""
    results: list[tuple[str, str | None, str | None]] = []
    pos = 0
    for m in _COLOR_SPAN_RE.finditer(text):
        if m.start() > pos:
            results.append((text[pos:m.start()], None, None))
        style_str = m.group(1)
        inner = m.group(2)
        color, bg = _parse_style(style_str)
        results.append((inner, color, bg))
        pos = m.end()
    if pos < len(text):
        results.append((text[pos:], None, None))
    if not results:
        results.append((text, None, None))
    return results


def _parse_style(style_str: str) -> tuple[str | None, str | None]:
    """Extract color and background-color from CSS style string."""
    text_color = None
    bg_color = None
    color_match = re.search(r"color:\s*([#\w]+)", style_str)
    if color_match:
        text_color = color_match.group(1)
    bg_match = re.search(r"background(?:-color)?:\s*([#\w]+)", style_str)
    if bg_match:
        bg_color = bg_match.group(1)
    return text_color, bg_color


def _parse_color(css_color: str) -> RGBColor | None:
    """Parse CSS color to RGBColor."""
    if css_color.startswith("#"):
        css_color = css_color.lstrip("#")
        if len(css_color) == 6:
            return RGBColor(
                int(css_color[0:2], 16),
                int(css_color[2:4], 16),
                int(css_color[4:6], 16),
            )
    return None


def _add_text_with_formatting(p, text: str, color: str | None, bg_color: str | None) -> None:
    """Add text to paragraph with **bold**, *italic*, and optional colors."""
    segments = _BOLD_RE.split(text)
    for i, seg in enumerate(segments):
        if i % 2 == 1:
            _add_segment(p, seg, bold=True, color=color, bg_color=bg_color)
        else:
            italic_segments = _ITALIC_RE.split(seg)
            for j, iseg in enumerate(italic_segments):
                if j % 2 == 1:
                    _add_segment(p, iseg, italic=True, color=color, bg_color=bg_color)
                else:
                    _add_segment(p, iseg, color=color, bg_color=bg_color)


def _add_segment(
    p, text: str, bold: bool = False, italic: bool = False,
    color: str | None = None, bg_color: str | None = None,
) -> None:
    """Add a single run with specified formatting."""
    if not text:
        return
    run = p.add_run(text)
    run.font.name = FONT_BODY
    run.bold = bold
    run.italic = italic
    if color:
        rgb = _parse_color(color)
        if rgb:
            run.font.color.rgb = rgb
    if bg_color:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), bg_color.lstrip("#"))
        shd.set(qn("w:val"), "clear")
        run._element.get_or_add_rPr().append(shd)


# ── Table support ──


def _add_table(doc: Document, table_lines: list[str]) -> None:
    """Parse pipe-table lines and add a docx table."""
    if len(table_lines) < 1:
        return

    rows: list[list[str]] = []
    for line in table_lines:
        cells = [c.strip() for c in line.strip("|").split("|")]
        rows.append(cells)

    if not rows:
        return

    max_cols = max(len(r) for r in rows)
    for r in rows:
        while len(r) < max_cols:
            r.append("")

    num_rows = len(rows)
    table = doc.add_table(rows=num_rows, cols=max_cols, style="Light Grid Accent 1")

    for i, row_data in enumerate(rows):
        for j, cell_text in enumerate(row_data):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            _add_inline_runs(p, cell_text)
            if i == 0:
                for run in p.runs:
                    run.bold = True
