"""Eval: financial_v2 确定性科目映射（A4 mapping）。

用法: python -m evals.test_financial_v2_mapping

覆盖（§5.7）：
- 文本规范化（Unicode/全角/空白/标点/编号，中文编号不误伤「一年内到期…」）；
- 内置规则集（报表类型内精确别名，无模糊包含「现金」≠「货币资金」）；
- 纯 map_candidate（精确命中 / 报表类型错配 / 无命中 / 空文本 / 无报表类型 /
  歧义多规则 / 排除词）；
- 端到端 map_record_set（落盘规则 + MAPPING_REQUIRED 问题 + 幂等不重复）。

全部合成 fixture（公司无关），临时 DB / 临时目录注入，不污染生产库。
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl

from financial_v2 import mapping
from financial_v2 import schema as S
from financial_v2 import store
from financial_v2 import excel_extractor as ex


def _tmp_db() -> str:
    fd, p = tempfile.mkstemp(suffix=".db", prefix="eval_fv2_map_")
    os.close(fd)
    return p


def _cleanup_db(p: str) -> None:
    for suffix in ("", "-wal", "-shm", "-journal"):
        try:
            os.remove(p + suffix)
        except FileNotFoundError:
            pass


def _tmp_xlsx(prefix: str = "eval_map_") -> str:
    fd, p = tempfile.mkstemp(suffix=".xlsx", prefix=prefix)
    os.close(fd)
    return p


def _cell(candidate_id: str, raw_item_text: str,
          statement_type_candidate: str | None) -> S.ExtractedFinancialCell:
    return S.ExtractedFinancialCell(
        candidate_id=candidate_id,
        record_set_version="rs-test",
        company_id="ACME",
        source_version="sv-test",
        statement_type_candidate=statement_type_candidate,
        raw_item_text=raw_item_text,
        raw_value_text=None,
        parsed_numeric_value=None,
        formula_text=None,
        cached_formula_value=None,
        period_text=None,
        period_candidate=None,
        period_type_candidate=None,
        scope_candidate=None,
        currency_candidate=None,
        unit_candidate=None,
        restatement_candidate=None,
        min_display_increment=None,
        locator=None,
        detection_evidence={},
        status="EXTRACTED",
        quality_flags=[],
        created_at="2026-01-01T00:00:00Z",
    )


def _build_workbook(path: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "资产负债表"
    ws.append(["资产负债表(单位：万元)", "2024-12-31"])
    ws.append(["货币资金", 1000])
    ws.append(["存货", 500])
    ws.append(["一年内到期的非流动资产", 300])
    ws.append(["其他非常规项目", 99])

    ws2 = wb.create_sheet("利润表")
    ws2.append(["利润表(单位：万元)", "2024-12-31"])
    ws2.append(["营业收入", 50000])
    ws2.append(["净利润", 8000])

    ws3 = wb.create_sheet("现金流量表")
    ws3.append(["现金流量表(单位：万元)", "2024-12-31"])
    ws3.append(["经营活动产生的现金流量净额", 12000])
    wb.save(path)


def _register_and_extract(db: str, xlsx: str) -> tuple[str, ex.ExcelExtractionResult]:
    store.init_db(db)
    file_hash = ex.sha256_file(xlsx)
    source_document_id = S.scope_source_document_id("ACME", "doc-map")
    source_version = S.derive_source_version(source_document_id, file_hash)
    store.register_source_atomic(
        S.FinancialSourceDocument(
            source_document_id=source_document_id, company_id="ACME",
            source_name="map.xlsx", source_class="financial_statement",
            declared_company_name="Acme", detected_company_name="Acme",
            subject_match_status="matched", created_at="2026-01-01T00:00:00Z"),
        S.FinancialSourceVersion(
            source_version=source_version, source_document_id=source_document_id,
            file_sha256=file_hash, file_type="xlsx", file_size=Path(xlsx).stat().st_size,
            document_id=None, document_version=None, created_at="2026-01-01T00:00:00Z"))
    result = ex.extract_excel(
        source_version,
        ex.ExcelExtractionPolicy(file_path=xlsx, dependency_versions={"openpyxl": openpyxl.__version__}),
        persist=True)
    return source_version, result


def main() -> dict:
    passed = 0
    failed = 0
    skipped = 0
    details: list[str] = []

    def check(cond, msg):
        nonlocal passed, failed
        if cond:
            passed += 1
            details.append(f"PASS: {msg}")
        else:
            failed += 1
            details.append(f"FAIL: {msg}")

    # ---- 文本规范化 ----
    check(mapping.normalize_item_text("1、资产总计") == "资产总计", "数字编号剥离")
    check(mapping.normalize_item_text("（一）货币资金") == "货币资金", "中文编号（括号）剥离")
    check(mapping.normalize_item_text("(1) 货币资金") == "货币资金", "半角括号编号剥离")
    check(mapping.normalize_item_text("一、资产总计") == "资产总计", "中文数字+顿号编号剥离")
    check(mapping.normalize_item_text("一年内到期的非流动资产") == "一年内到期的非流动资产",
          "中文数字不作为编号误伤")
    check(mapping.normalize_item_text(" 货币 资金 ") == "货币资金", "空白剥离")
    check(mapping.normalize_item_text("货币资金（净额）") == "货币资金净额", "标点剥离（括号内保留）")
    check(mapping.normalize_item_text("") == "", "空串 → 空串")
    check(mapping.normalize_item_text("ＣＡＳＨ") == "cash", "全角字母转半角")

    # ---- 内置规则集 ----
    rules = mapping.build_builtin_rules()
    check(len(rules) == 136, f"内置规则 136 条（实际 {len(rules)}）")
    check(len({r.rule_id for r in rules}) == 136, "rule_id 唯一")
    by_type = {"balance_sheet": 0, "income_statement": 0, "cash_flow": 0}
    for r in rules:
        by_type[r.statement_type] += 1
        check(r.standard_item_code == r.rule_id, f"rule_id 与 standard_item_code 一致（{r.rule_id}）")
    check(by_type["balance_sheet"] == 65, f"资产负债表规则 65（实际 {by_type['balance_sheet']}）")
    check(by_type["income_statement"] == 34, f"利润表规则 34（实际 {by_type['income_statement']}）")
    check(by_type["cash_flow"] == 37, f"现金流量表规则 37（实际 {by_type['cash_flow']}）")

    # ---- 纯 map_candidate ----
    check(mapping.map_candidate(_cell("c1", "货币资金", "balance_sheet"), rules).standard_item_code
          == "CASH_AND_EQUIVALENTS", "货币资金 → CASH_AND_EQUIVALENTS")
    check(mapping.map_candidate(_cell("c2", "净利润", "income_statement"), rules).standard_item_code
          == "NET_PROFIT", "净利润 → NET_PROFIT")
    check(mapping.map_candidate(_cell("c3", "经营活动现金流量净额", "cash_flow"), rules).standard_item_code
          == "OPERATING_CASH_FLOW", "简称「经营活动现金流量净额」→ OPERATING_CASH_FLOW")
    check(mapping.map_candidate(_cell("c4", "经营活动产生的现金流量净额", "cash_flow"), rules).standard_item_code
          == "OPERATING_CASH_FLOW", "全称 → OPERATING_CASH_FLOW")

    # 报表类型错配：货币资金 放在利润表里不得命中资产负债表规则。
    out = mapping.map_candidate(_cell("c5", "货币资金", "income_statement"), rules)
    check(out.status == "mapping_required" and out.reason == "no_rule_match",
          "报表类型错配 → no_rule_match（不跨表）")

    # 无命中 / 空文本 / 无报表类型。
    out = mapping.map_candidate(_cell("c6", "某个未知科目", "balance_sheet"), rules)
    check(out.status == "mapping_required" and out.reason == "no_rule_match", "未知科目 → no_rule_match")
    out = mapping.map_candidate(_cell("c7", "", "balance_sheet"), rules)
    check(out.status == "mapping_required" and out.reason == "empty_item_text", "空文本 → empty_item_text")
    out = mapping.map_candidate(_cell("c8", "货币资金", None), rules)
    check(out.status == "mapping_required" and out.reason == "no_statement_type",
          "无报表类型 → no_statement_type")

    # 模糊包含防线：简称「现金」不得命中「货币资金」。
    out = mapping.map_candidate(_cell("c9", "现金", "balance_sheet"), rules)
    check(out.status == "mapping_required" and out.reason == "no_rule_match",
          "「现金」≠「货币资金」（不做模糊包含）")

    # 歧义：注入同报表类型、同别名、不同 code 的两条规则 → ambiguous_mapping。
    dup_rules = [
        S.MappingRule("R_A", "t", "balance_sheet", "X_CODE_A", ["某科目"], [], 100, "2026-01-01"),
        S.MappingRule("R_B", "t", "balance_sheet", "X_CODE_B", ["某科目"], [], 100, "2026-01-01"),
    ]
    out = mapping.map_candidate(_cell("c10", "某科目", "balance_sheet"), dup_rules)
    check(out.status == "mapping_required" and out.reason == "ambiguous_mapping"
          and len(out.conflicting_rules) == 2, "歧义多规则 → ambiguous_mapping（列出冲突）")

    # 排除词：规则别名命中但排除词是原文子串 → 整条规则失效。
    excl_rules = [
        S.MappingRule("R_C", "t", "balance_sheet", "X_CODE_C", ["某科目"], ["排除"], 100, "2026-01-01"),
    ]
    out = mapping.map_candidate(_cell("c11", "某科目排除", "balance_sheet"), excl_rules)
    check(out.status == "mapping_required" and out.reason == "no_rule_match",
          "排除词命中 → 规则失效")

    # ---- 端到端 map_record_set ----
    db = _tmp_db()
    xlsx = _tmp_xlsx()
    try:
        _build_workbook(xlsx)
        _, extract = _register_and_extract(db, xlsx)
        rs = extract.record_set_version

        # 首轮映射：落盘规则 + MAPPING_REQUIRED。
        result = mapping.map_record_set(rs, persist=True)
        mapped_codes = {o.standard_item_code for o in result.outcomes if o.status == "mapped"}
        check("CASH_AND_EQUIVALENTS" in mapped_codes, "货币资金 已映射")
        check("OPERATING_REVENUE" in mapped_codes, "营业收入 已映射")
        check("NET_PROFIT" in mapped_codes, "净利润 已映射")
        check("OPERATING_CASH_FLOW" in mapped_codes, "经营现金流 已映射")
        check("NON_CURRENT_ASSET_DUE_WITHIN_ONE_YEAR" in mapped_codes, "一年内到期非流动资产 已映射")

        unmapped = [o for o in result.outcomes if o.status == "mapping_required"]
        check(len(unmapped) == 1 and unmapped[0].raw_item_text == "其他非常规项目",
              f"未知科目「其他非常规项目」→ MAPPING_REQUIRED（实际 {len(unmapped)} 条）")
        check(result.issues_committed == 1, f"落盘 1 条 MAPPING_REQUIRED（实际 {result.issues_committed}）")

        # 规则已持久化。
        persisted = store.list_mapping_rules(rule_version="1.0")
        check(len(persisted) == 136, f"内置规则已落盘 136 条（实际 {len(persisted)}）")

        # 问题落库且引用正确。
        issues = store.list_extraction_issues(rs)
        map_issues = [i for i in issues if i.issue_type == "MAPPING_REQUIRED"]
        check(len(map_issues) == 1, "extraction_issue 表有 1 条 MAPPING_REQUIRED")
        check(map_issues[0].candidate_id == unmapped[0].candidate_id, "问题引用正确候选")

        # 幂等：二次映射不重复落盘、计数一致。
        result2 = mapping.map_record_set(rs, persist=True)
        check(result2.mapped_count == result.mapped_count
              and result2.unmapped_count == result.unmapped_count, "二次映射计数一致")
        check(result2.issues_committed == 0, "二次映射不重复写问题")
        check(len(store.list_extraction_issues(rs)) == 1, "问题未重复（仍 1 条）")

        # validate-only：不落盘规则/问题。
        db2 = _tmp_db()
        xlsx2 = _tmp_xlsx()
        try:
            _build_workbook(xlsx2)
            _, extract2 = _register_and_extract(db2, xlsx2)
            rs2 = extract2.record_set_version
            rv = mapping.map_record_set(rs2, persist=False)
            check(rv.issues_committed == 0, "validate-only 不写问题")
            check(len(store.list_mapping_rules(rule_version="1.0")) == 0,
                  "validate-only 不写规则")
            check(rv.mapped_count == 6, f"validate-only 仍产出映射结果（实际 {rv.mapped_count}）")
        finally:
            _cleanup_db(db2)
            try:
                os.remove(xlsx2)
            except FileNotFoundError:
                pass
    finally:
        _cleanup_db(db)
        try:
            os.remove(xlsx)
        except FileNotFoundError:
            pass

    return {"passed": passed, "failed": failed, "skipped": skipped, "details": details}


if __name__ == "__main__":
    import json

    result = main()
    print(json.dumps(result, ensure_ascii=False, indent=2))
    sys.exit(0 if result["failed"] == 0 else 1)
