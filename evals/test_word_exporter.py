"""Eval: Word export — Markdown → docx。

用法: python -m evals.test_word_exporter
"""

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

os = __import__("os")


def main() -> dict:
    passed = 0
    failed = 0
    skipped = 0
    details: list[str] = []

    def check(condition, msg):
        nonlocal passed, failed
        if condition:
            passed += 1
            details.append(f"PASS: {msg}")
        else:
            failed += 1
            details.append(f"FAIL: {msg}")

    try:
        from docx import Document
    except ImportError:
        return {"passed": 0, "failed": 0, "skipped": 1,
                "details": ["SKIP: python-docx not installed"]}

    from reporting.word_exporter import export

    # ── Test 1: Basic heading + paragraph ──
    md = "# 测试标题\n\n这是正文段落。\n\n## 二级标题\n\n第二段内容。"
    fd, tmp = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    try:
        export(md, tmp)
        doc = Document(tmp)
        check(len(doc.paragraphs) >= 4, f"Basic: has paragraphs (got {len(doc.paragraphs)})")
        full_text = "\n".join(p.text for p in doc.paragraphs)
        check("测试标题" in full_text, "Basic: H1 preserved")
        check("二级标题" in full_text, "Basic: H2 preserved")
        check("正文段落" in full_text, "Basic: body text preserved")
    finally:
        Path(tmp).unlink(missing_ok=True)

    # ── Test 2: Bold and italic ──
    md = "# 格式测试\n\n这是**粗体**文字和*斜体*文字。"
    fd, tmp = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    try:
        export(md, tmp)
        doc = Document(tmp)
        runs = []
        for p in doc.paragraphs:
            runs.extend(p.runs)
        check(len(runs) >= 3, f"Format: has inline runs (got {len(runs)})")
        bold_runs = [r for r in runs if r.bold]
        check(len(bold_runs) >= 1, "Format: bold detected")
        italic_runs = [r for r in runs if r.italic]
        check(len(italic_runs) >= 1, "Format: italic detected")
    finally:
        Path(tmp).unlink(missing_ok=True)

    # ── Test 3: Table ──
    md = "# 表格测试\n\n| 指标 | 数值 |\n|------|------|\n| 流动比率 | 1.60 |\n| 速动比率 | 1.20 |"
    fd, tmp = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    try:
        export(md, tmp)
        doc = Document(tmp)
        check(len(doc.tables) >= 1, f"Table: has table (got {len(doc.tables)})")
        if doc.tables:
            table = doc.tables[0]
            check(len(table.rows) == 3, f"Table: 3 rows (got {len(table.rows)})")
            check("流动比率" in table.cell(1, 0).text, "Table: cell content preserved")
    finally:
        Path(tmp).unlink(missing_ok=True)

    # ── Test 4: Lists ──
    md = "# 列表测试\n\n- 项目一\n- 项目二\n\n1. 第一\n2. 第二"
    fd, tmp = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    try:
        export(md, tmp)
        doc = Document(tmp)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        check("项目一" in full_text, "List: unordered item preserved")
        check("第一" in full_text, "List: ordered item preserved")
    finally:
        Path(tmp).unlink(missing_ok=True)

    # ── Test 5: Empty input ──
    fd, tmp = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    try:
        export("", tmp)
        doc = Document(tmp)
        check(len(doc.paragraphs) <= 2, f"Empty: minimal output (got {len(doc.paragraphs)})")
    finally:
        Path(tmp).unlink(missing_ok=True)

    # ── Test 6: Verifier colored span ──
    md = '# 回检\n\n资产负债率降至<span style="background:#fff3cd" title="数值差异5.3%">61.94%</span>。'
    fd, tmp = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    try:
        export(md, tmp)
        doc = Document(tmp)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        check("61.94%" in full_text, "Span: inner text preserved")
        check("数值差异" not in full_text, "Span: title attribute excluded from text")
    finally:
        Path(tmp).unlink(missing_ok=True)

    return {"passed": passed, "failed": failed, "skipped": skipped, "details": details}


if __name__ == "__main__":
    import json as _json
    result = main()
    for d in result["details"]:
        print(d)
    print(f"\n{result['passed']} passed, {result['failed']} failed, "
          f"{result['skipped']} skipped")
    sys.exit(0 if result["failed"] == 0 else 1)
